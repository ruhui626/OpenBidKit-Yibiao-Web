"""Small synthetic image, DOCX and image-only PDF inputs for offline OCR verification."""
import argparse
import json
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader, PdfWriter
from pypdf.generic import RectangleObject
from reportlab.pdfgen.canvas import Canvas


def main():
    p = argparse.ArgumentParser(); p.add_argument('--output', type=Path, required=True)
    out = p.parse_args().output.resolve(); out.mkdir(parents=True, exist_ok=False)
    size = (1800, 2400)
    image = Image.new('RGB', size, 'white'); draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(r'C:\Windows\Fonts\msyh.ttc', 58)
    lines = ['模拟样本  非真实证书', '项目名称：青松路道路维修', '项目经理：张明',
             '证书编号：AB0123456789', '有效期：2028年12月31日', '投标报价：1234567.89元']
    regions = []
    for i, text in enumerate(lines):
        xy = (150, 180 + i * 260)
        box = draw.textbbox(xy, text, font=font)
        draw.text(xy, text, font=font, fill=(15, 15, 15))
        regions.append({'text': text, 'rect': [box[0], box[1], box[2]-box[0], box[3]-box[1]]})
    image.save(out / 'clean.png')
    image.rotate(5, resample=Image.Resampling.BICUBIC, fillcolor='white').save(out / 'skew.png')
    degraded = image.resize((540, 720), Image.Resampling.BILINEAR).filter(ImageFilter.GaussianBlur(0.85)).resize(size, Image.Resampling.BILINEAR)
    stamp = ImageDraw.Draw(degraded); stamp.ellipse((700, 650, 1190, 1140), outline=(210, 45, 45), width=17)
    degraded.save(out / 'degraded.jpg', quality=28)
    unreadable = image.resize((90, 120), Image.Resampling.BILINEAR).filter(ImageFilter.GaussianBlur(1.8)).resize(size, Image.Resampling.BILINEAR)
    unreadable.save(out / 'unreadable.png')
    Image.new('RGB', size, 'white').save(out / 'blank.png')

    doc = Document(); doc.add_heading('图片文字提取验证样本', 0)
    doc.add_paragraph('以下均为人工构造的模拟资料。正文未写入图片中的项目、人员和编号。')
    doc.add_picture(str(out / 'clean.png'), width=Inches(3))
    doc.add_paragraph('同一图片再次使用，需要区分两个位置。')
    doc.add_picture(str(out / 'clean.png'), width=Inches(2))
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].add_run().add_picture(str(out / 'degraded.jpg'), width=Inches(3))
    doc.sections[0].header.paragraphs[0].add_run().add_picture(str(out / 'clean.png'), width=Inches(1))
    doc.sections[0].footer.paragraphs[0].add_run().add_picture(str(out / 'degraded.jpg'), width=Inches(1))
    doc.add_paragraph('裁剪及旋转图片只验证原图识别，不能据此推断页面显示范围。')
    shape = doc.add_picture(str(out / 'clean.png'), width=Inches(2))
    blip_fill = shape._inline.find('.//' + qn('pic:blipFill'))
    src_rect = OxmlElement('a:srcRect'); src_rect.set('l', '30000'); blip_fill.insert(1, src_rect)
    transform = shape._inline.find('.//' + qn('a:xfrm')); transform.set('rot', '5400000')
    doc.save(out / 'images.docx')

    # Only images are painted; PDF text extraction must remain empty on every page.
    pdf = Canvas(str(out / 'scan-base.pdf'), pagesize=(540, 720))
    for name in ['clean.png', 'degraded.jpg', 'clean.png']:
        pdf.drawImage(str(out / name), 0, 0, width=540, height=720)
        pdf.showPage()
    pdf.save()
    reader = PdfReader(out / 'scan-base.pdf'); writer = PdfWriter()
    for i, page in enumerate(reader.pages):
        if i == 2:
            page.cropbox = RectangleObject([15, 12, 525, 708]); page.rotate(90)
        writer.add_page(page)
    with (out / 'scans.pdf').open('wb') as stream: writer.write(stream)
    manifest = {'synthetic': True, 'imageSize': list(size), 'regions': regions,
                'standaloneImages': ['clean.png', 'degraded.jpg', 'skew.png', 'unreadable.png', 'blank.png'],
                'docx': {'name': 'images.docx', 'occurrences': 6, 'uniqueMedia': 2, 'croppedRotatedOccurrences': 1},
                'pdf': {'name': 'scans.pdf', 'pages': 3, 'renderScale': 2, 'manualUprightCCWByPage': {'3': 90},
                        'displayCoordinates': 'top-left of cropped/rotated displayed page, in points; not raw PDF user coordinates'},
                'groundTruthIsNotOcr': True}
    (out / 'fixture-manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'fixtures': str(out), 'images': 5, 'docx': 1, 'imageOnlyPdfPages': 3}))


if __name__ == '__main__': main()
