"""Author DOCX pairs, text-only opinions and separate human answers. No analyzer import."""
import argparse
import json
import sys
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageFont

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'word-evidence-spike'))
from make_textbox_fixtures import Builder


def make_doc(title):
    doc=Document();doc.add_heading(title,0)
    return doc


def workflow(path, after):
    doc=make_doc('道路维修投标文件模拟稿')
    project='演示新路维修工程' if after else '演示旧路维修工程'
    paras=[f'本工程项目名称：{project}。',f'投标函所列项目名称：{project}。',
           '施工总工期：30天。',f'夜间施工联系电话：{"13900000000" if after else "13800000000"}。',
           f'人员安排：项目经理{"赵乙" if after else "赵甲"}负责现场管理。',
           '材料承诺：水泥编号为'+('批次新号' if after else '批次旧号')+'，砂石编号为批次旧号。',
           '资料清单：授权委托书待补充。','施工方案：道路分段封闭后维修。']
    if after:
        doc.add_paragraph('新增编制说明：本稿仅供实验使用。')
        table=doc.add_table(rows=1,cols=2)
        table.cell(0,0).text='递交记录';table.cell(0,1).text='模拟记录'
        paras=[paras[i] for i in [4,2,0,6,1,7,5,3]]
    for text in paras:doc.add_paragraph(text)
    if not after:doc.add_paragraph('质保承诺：质保期为12个月。')
    table=doc.add_table(rows=1,cols=3)
    for cell,text in zip(table.rows[0].cells,['岗位','姓名','说明']):cell.text=text
    rows=[['项目经理','赵甲','现场负责人'],['安全员','李丙','现场巡查']]
    if after:rows=[['资料员','黄丁','归档'],rows[1],rows[0]]
    for row in rows:
        for cell,text in zip(table.add_row().cells,row):cell.text=text
    doc.sections[0].header.paragraphs[0].text=f'工程标识：{project}'
    doc.sections[0].footer.paragraphs[0].text='归档编号：备案旧号A01。'
    box=Builder()
    box.box(doc.add_paragraph('投标联系人信息框').add_run()._r,
            [f'投标联系人：{"周乙" if after else "周甲"}。'],mode='legacy')
    doc.save(path)


def ambiguity(path,after):
    doc=make_doc('位置及重复内容模拟稿')
    values=[f'采购承诺：合同编号为{"合同新号" if after else "合同旧号"}。',
            f'工程别名：{"新编号乙" if after else "旧编号甲"}。',
            '安全员：赵甲。',f'项目经理：{"赵乙" if after else "赵甲"}。']
    if after:values=[values[3],values[1],values[0],values[2],values[0],
                    '运输方案：道路运输集中在夜间。','现场设置专人指挥。',
                    '补充附件仍引用旧编号甲。']
    else:values.append('运输方案：道路运输集中在白天，安排专人指挥。')
    values += [f'重复声明：代理人{"孙乙" if after else "孙甲"}。']*2
    values += ['重复承诺：负责人'+('郑乙' if after else '郑甲')+'。','重复承诺：负责人郑甲。']
    for text in values:doc.add_paragraph(text)
    doc.save(path)


def coverage(path,after,out):
    project='模拟新校区工程' if after else '模拟旧校区工程'
    image=Image.new('RGB',(1200,600),'white');draw=ImageDraw.Draw(image)
    font=ImageFont.truetype(r'C:\Windows\Fonts\msyh.ttc',48)
    draw.text((45,65),'模拟图片 非真实材料',font=font,fill='black')
    draw.text((45,200),project,font=font,fill='black')
    draw.text((45,335),'证书编号：DE0123456789',font=font,fill='black')
    image_path=out/('scan-after.png' if after else 'scan-before.png');image.save(image_path)
    doc=make_doc('带扫描附件的模拟稿')
    doc.add_paragraph(f'本项目名称：{project}。')
    doc.add_paragraph('证书编号请核对下面扫描图片。')
    doc.add_picture(str(image_path))
    doc.add_paragraph('施工方案：施工期间采用分段围挡。')
    doc.save(path)


def vocabulary(path,after):
    doc=make_doc('桥梁修复投标模拟稿')
    project='东桥修复工程' if after else '南桥修复工程'
    doc.sections[0].header.paragraphs[0].text='合同项目：南桥修复工程。'
    doc.add_paragraph('参考方案曾提到75天，与本次维修工期无关。')
    if after:doc.add_paragraph('新增变更记录：段落及表格顺序调整。')
    texts=[f'合同项目：{project}。','维修工期：60天。']
    if after:texts.reverse()
    for text in texts:doc.add_paragraph(text)
    table=doc.add_table(rows=1,cols=2)
    table.cell(0,0).text='岗位';table.cell(0,1).text='人员'
    for label,person in ([('安全员','许三'),('技术负责人','林二')] if after else [('技术负责人','陈一'),('安全员','许三')]):
        row=table.add_row();row.cells[0].text=label;row.cells[1].text=person
    doc.save(path)


def main():
    p=argparse.ArgumentParser();p.add_argument('--output',type=Path,required=True)
    out=p.parse_args().output.resolve();out.mkdir(parents=True,exist_ok=False)
    cases=[];answers={}
    def case(name,builder,entries):
        before=name+'-before.docx';after=name+'-after.docx'
        builder(out/before,False);builder(out/after,True)
        opinions=[]
        for number,(text,truth,decision,note) in enumerate(entries,1):
            issue_id=f'{name}-{number:02d}'
            opinions.append({'id':issue_id,'text':text})
            answers[issue_id]={'truthStatus':truth,'expectedConservativeDecision':decision,'authorNote':note}
        cases.append({'id':name,'before':before,'after':after,'opinions':opinions})
    case('workflow',workflow,[
        ('把全文的“演示旧路维修工程”改为“演示新路维修工程”。','modified','modified','正文两处和页眉均改正，段落顺序改变。'),
        ('把“项目经理”相关位置的“赵甲”改为“赵乙”。','partially_modified','partially_modified','正文改，人员表未改；人员表位置及行顺序改变。'),
        ('把“施工总工期”相关位置的“30天”改为“45天”。','unmodified','unmodified','实际仍为30天。'),
        ('把“夜间施工联系电话”相关位置的“13800000000”改为“13900000000”。','modified','modified','整段移动后号码改正。'),
        ('把“投标联系人”相关位置的“周甲”改为“周乙”。','modified','modified','只在文本框内修改。'),
        ('把“归档编号”相关位置的“备案旧号A01”改为“备案新号A02”。','unmodified','unmodified','页脚保留原编号。'),
        ('把“质保承诺”相关位置的“12个月”改为“24个月”。','unmodified','undetermined','原段被删，未提供24个月承诺；原型不得凭旧值消失报已改。'),
        ('把“材料承诺”相关位置的“批次旧号”改为“批次新号”。','partially_modified','partially_modified','同一段落两处只改一处。'),
        ('请补齐授权委托书并盖章。','undetermined','undetermined','无可校验盖章材料或可执行文本替换目标。'),
        ('请优化整体施工方案。','undetermined','undetermined','意见缺乏可验证标准。')])
    case('ambiguity',ambiguity,[
        ('把“采购承诺”相关位置的“合同旧号”改为“合同新号”。','modified','undetermined','已改并新增同文副本，数量变化，原型保守拒绝唯一对应。'),
        ('把“运输方案”相关位置的“白天”改为“夜间”。','modified','undetermined','已改并拆成两个段落；原型未验证语义拆分。'),
        ('把全文的“旧编号甲”改为“新编号乙”。','partially_modified','partially_modified','原段改正，新增加的附件仍有旧号。'),
        ('把“项目经理”相关位置的“赵甲”改为“赵乙”。','modified','modified','同名安全员未改，属于本意见范围以外。'),
        ('把“重复声明”相关位置的“孙甲”改为“孙乙”。','modified','modified','两份相同声明均改，按组核对，不编造逐段血缘。'),
        ('把“重复承诺”相关位置的“郑甲”改为“郑乙”。','partially_modified','partially_modified','两份相同承诺只改一份。')])
    case('coverage',lambda path,after:coverage(path,after,out),[
        ('把全文的“模拟旧校区工程”改为“模拟新校区工程”。','modified','undetermined','正文和图片实际都改了，但本轮不联动OCR，全文覆盖不足。'),
        ('把正文可编辑文字的“模拟旧校区工程”改为“模拟新校区工程”。','modified','modified','意见明确限于可编辑正文，不承诺扫描图。'),
        ('请核对扫描件中的证书编号是否已经更正。','undetermined','undetermined','缺乏指定旧值和目标新值，本轮不作OCR判断。'),
        ('请确认施工方案满足夜间作业和交通导改要求。','undetermined','undetermined','没有专业验收标准和招标条款。')])
    case('vocabulary',vocabulary,[
        ('把正文可编辑文字的“南桥修复工程”改为“东桥修复工程”。','modified','modified','改换词汇后仍能对应移动段落，页眉不在本意见范围。'),
        ('把全文的“南桥修复工程”改为“东桥修复工程”。','partially_modified','partially_modified','正文已改但页眉未改。'),
        ('把“技术负责人”相关位置的“陈一”改为“林二”。','modified','modified','不同人名和岗位，表格行位置改变。'),
        ('把“维修工期”相关位置的“60天”改为“75天”。','unmodified','unmodified','其他参考段有75天，实际工期段仍为60天。')])
    (out/'cases.json').write_text(json.dumps({'synthetic':True,'cases':cases},ensure_ascii=False,indent=2),encoding='utf-8')
    (out/'human-answers.json').write_text(json.dumps(answers,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'pairs':len(cases),'opinions':len(answers),'output':str(out)},ensure_ascii=False))


if __name__=='__main__':main()
