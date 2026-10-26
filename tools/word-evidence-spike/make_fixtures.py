"""Create synthetic DOCX fixtures, never read or alter a customer document."""
import argparse
import io
import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from lxml import etree


def setup(title):
    doc = Document()
    doc.core_properties.author = "Synthetic test fixture"
    doc.core_properties.last_modified_by = "Synthetic test fixture"
    doc.core_properties.title = title
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.8)
    for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_paragraph(title, "Title")
    section.header.paragraphs[0].text = "页眉旧项目：演示旧路维修工程"
    section.footer.paragraphs[0].text = "页脚：虚构企业，仅供定位实验"
    return doc


def fixture_pair(out):
    manifest = {}
    for mode in ["before", "after"]:
        after = mode == "after"
        doc = setup("正文与表格定位验证")
        doc.add_heading("一 项目概况", 1)
        if after:
            doc.add_paragraph("新增说明：段落顺序变化用于检验版本绑定。")
        p = doc.add_paragraph()
        p.add_run("项目名称：")
        p.add_run("演示新路" if after else "演示旧路").bold = True
        p.add_run("维修工程。")
        manager = "示例乙" if after else "示例甲"
        doc.add_paragraph(f"项目经理：{manager}。")
        doc.add_paragraph("材料要求：A&B <C>，不可遗漏。")
        p = doc.add_paragraph("计划：")
        p.add_run().add_tab()
        p.add_run("施工准备")
        p.add_run().add_break()
        p.add_run("恢复交通")
        doc.add_paragraph("复核备注：待核对。")
        doc.add_heading("二 人员与组织", 1)
        table = doc.add_table(rows=4, cols=3)
        table.style = "Table Grid"
        for row, values in enumerate([
            ["岗位", "姓名", "说明"], ["项目经理", "示例甲", "负责人"],
            ["安全员", "示例丙", "现场检查"], ["合并单元格：进场要求", "", "到岗"],
        ]):
            for column, value in enumerate(values):
                table.cell(row, column).text = value
        table.cell(3, 0).merge(table.cell(3, 1))
        nested = table.cell(2, 2).add_table(rows=2, cols=2)
        nested.style = "Table Grid"
        for row, values in enumerate([["检查项", "状态"], ["隐患记录", "复核备注：待核对。"]]):
            for col, value in enumerate(values):
                nested.cell(row, col).text = value
        doc.add_paragraph("复核备注：待核对。")
        doc.add_heading("三 阶段安排", 1)
        vertical = doc.add_table(rows=3, cols=2)
        vertical.style = "Table Grid"
        for row, values in enumerate([["阶段", "任务"], ["第一阶段", "施工准备"], ["", "恢复交通"]]):
            for col, value in enumerate(values):
                vertical.cell(row, col).text = value
        vertical.cell(1, 0).merge(vertical.cell(2, 0))
        filename = f"{mode}.docx"
        doc.save(out / filename)
        manifest[filename] = {
            "source": "synthetic", "expectedTableCount": 3,
            "expectations": [
                {"text": f"项目名称：{'演示新路' if after else '演示旧路'}维修工程。", "kind": "body-paragraph"},
                {"text": f"项目经理：{manager}。", "kind": "body-paragraph"},
                {"text": "材料要求：A&B <C>，不可遗漏。", "kind": "body-paragraph"},
                {"text": "计划：\t施工准备\n恢复交通", "kind": "body-paragraph"},
                {"text": "示例甲", "kind": "table-paragraph", "tableIndex": 1, "xmlRow": 2, "gridColumn": 2},
                {"text": "合并单元格：进场要求", "tableIndex": 1, "xmlRow": 4, "gridColumn": 1, "gridSpan": 2},
                {"text": "隐患记录", "tableIndex": 2, "xmlRow": 2, "gridColumn": 1, "tableDepth": 2},
                {"text": "第一阶段", "tableIndex": 3, "xmlRow": 2, "gridColumn": 1, "verticalMerge": "restart"},
                {"text": "恢复交通", "tableIndex": 3, "xmlRow": 3, "gridColumn": 2},
            ],
            "duplicateQuote": "复核备注：待核对。", "duplicateCount": 3,
            "mustNotInclude": ["页眉旧项目：演示旧路维修工程", "页脚：虚构企业，仅供定位实验"],
        }
    return manifest


def add_text(parent, text, tag="w:t"):
    node = OxmlElement(tag)
    node.text = text
    parent.append(node)
    return node


def fixture_edges(out):
    doc = setup("未覆盖内容边界验证")
    doc.add_heading("一 正文和特殊部件", 1)
    doc.add_paragraph("正文可见：计划工期45天。")
    doc.add_paragraph("现场🌧️记录：Unicode 偏移测试。")
    control = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    add_text(r, "内容控件：现场负责人示例丙。")
    p.append(r)
    content.append(p)
    control.append(content)
    doc.element.body.insert(len(doc.element.body) - 1, control)
    p = doc.add_paragraph("工期：")
    for tag, text, text_tag in [("w:del", "30天", "w:delText"), ("w:ins", "45天", "w:t")]:
        change = OxmlElement(tag)
        change.set(qn("w:id"), "1" if tag == "w:del" else "2")
        change.set(qn("w:author"), "Synthetic")
        change.set(qn("w:date"), "2026-10-24T00:00:00Z")
        run = OxmlElement("w:r")
        add_text(run, text, text_tag)
        change.append(run)
        p._p.append(change)
    p.add_run("。")
    p = doc.add_paragraph("文本框容器。")
    run = OxmlElement("w:r")
    pict = OxmlElement("w:pict")
    shape = etree.Element("{urn:schemas-microsoft-com:vml}shape", nsmap={"v": "urn:schemas-microsoft-com:vml"})
    shape.set("id", "synthetic_textbox")
    shape.set("style", "width:200pt;height:30pt")
    textbox = etree.Element("{urn:schemas-microsoft-com:vml}textbox")
    text_content = OxmlElement("w:txbxContent")
    inner_p, inner_r = OxmlElement("w:p"), OxmlElement("w:r")
    add_text(inner_r, "文本框内旧项目：演示旧路维修工程。")
    inner_p.append(inner_r)
    text_content.append(inner_p)
    textbox.append(text_content)
    shape.append(textbox)
    pict.append(shape)
    run.append(pict)
    p._p.append(run)
    image = Image.new("RGB", (600, 70), "white")
    ImageDraw.Draw(image).text((10, 12), "图片中的项目经理：示例甲", fill="black", font=ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc", 30))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)
    doc.add_picture(buffer, width=Inches(4))
    p = doc.add_paragraph("生成日期：")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "DATE")
    r = OxmlElement("w:r")
    add_text(r, "2026-10-24")
    field.append(r)
    p._p.append(field)
    p = doc.add_paragraph("自动列表编号不包含在纯文本中", "List Number")
    hidden = doc.add_paragraph().add_run("隐藏文字：演示旧路维修工程。")
    hidden.font.hidden = True
    footnote_xml = b'<?xml version="1.0" encoding="UTF-8"?><w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="2"><w:p><w:r><w:t>FOOTNOTE_ONLY_SYNTHETIC</w:t></w:r></w:p></w:footnote></w:footnotes>'
    part = Part(PackURI("/word/footnotes.xml"), "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml", footnote_xml, doc.part.package)
    doc.part.relate_to(part, RT.FOOTNOTES)
    p = doc.add_paragraph("脚注标记")
    r = OxmlElement("w:r")
    reference = OxmlElement("w:footnoteReference")
    reference.set(qn("w:id"), "2")
    r.append(reference)
    p._p.append(r)
    p = doc.add_paragraph()
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), doc.part.relate_to("https://example.invalid/evidence-test", RT.HYPERLINK, is_external=True))
    r = OxmlElement("w:r")
    add_text(r, "超链接文字：测试网址")
    hyperlink.append(r)
    p._p.append(hyperlink)
    doc.save(out / "edges.docx")
    return {"source": "synthetic", "expectedTableCount": 0, "expectations": [
        {"text": "正文可见：计划工期45天。"}, {"text": "内容控件：现场负责人示例丙。"},
        {"text": "工期：45天。", "requiredFlag": "tracked_changes_require_review"},
        {"text": "生成日期：2026-10-24", "requiredFlag": "cached_field_result_only"},
        {"text": "隐藏文字：演示旧路维修工程。", "requiredFlag": "hidden_text_requires_review"},
        {"text": "超链接文字：测试网址"}, {"text": "现场🌧️记录：Unicode 偏移测试。"},
    ], "mustNotInclude": ["文本框内旧项目：演示旧路维修工程。", "FOOTNOTE_ONLY_SYNTHETIC", "图片中的项目经理：示例甲"]}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    if any(args.output.iterdir()):
        raise SystemExit("Use a new empty fixture directory; no existing files will be overwritten.")
    manifest = fixture_pair(args.output)
    manifest["edges.docx"] = fixture_edges(args.output)
    (args.output / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"created": list(manifest), "synthetic": True}, ensure_ascii=False))


if __name__ == "__main__":
    main()
