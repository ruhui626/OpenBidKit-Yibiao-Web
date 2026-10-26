"""Synthetic multi-section fixtures with hand-authored binding expectations."""
import argparse
import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.shared import Inches, Pt, RGBColor
from lxml import etree


def new_document():
    document = Document()
    document.core_properties.author = document.core_properties.last_modified_by = 'Synthetic header/footer test'
    for name in ['Normal', 'Title', 'Heading 1', 'Header', 'Footer']:
        style = document.styles[name]
        style.font.name = 'Microsoft YaHei'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    document.styles['Normal'].font.size = Pt(10)
    document.add_paragraph('分节文档验证样本', 'Title')
    return document


def fill(container, text, table_label=None, page_field=False):
    container.is_linked_to_previous = False
    paragraph = container.paragraphs[0]
    cut = text.find('：') + 1
    paragraph.add_run(text[:cut])
    paragraph.add_run(text[cut:]).bold = True
    container.add_paragraph('重复标记：待核对。')
    if table_label:
        table = container.add_table(rows=1, cols=2, width=Inches(5))
        table.style = 'Table Grid'
        table.cell(0, 0).text = '表格字段'
        table.cell(0, 1).text = table_label
    if page_field:
        p = container.add_paragraph('页码：')
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), 'PAGE')
        run = OxmlElement('w:r')
        text_node = OxmlElement('w:t')
        text_node.text = '7'
        run.append(text_node)
        field.append(run)
        p._p.append(field)
    return str(container.part.partname).lstrip('/')


def write_package(path, document, rename=None):
    # Serialize once, then patch only this new in-memory package; no existing DOCX is opened for editing.
    import io
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    with ZipFile(buffer) as archive:
        parts = {n: archive.read(n) for n in archive.namelist()}
    if rename:
        old, new = rename
        parts[new] = parts.pop(old)
        rel = etree.fromstring(parts['word/_rels/document.xml.rels'])
        for node in rel:
            if node.get('Target') == old.removeprefix('word/'):
                node.set('Target', new.removeprefix('word/'))
        parts['word/_rels/document.xml.rels'] = etree.tostring(rel, xml_declaration=True, encoding='UTF-8')
        types = etree.fromstring(parts['[Content_Types].xml'])
        for node in types:
            if node.get('PartName') == '/' + old:
                node.set('PartName', '/' + new)
        parts['[Content_Types].xml'] = etree.tostring(types, xml_declaration=True, encoding='UTF-8')
    with ZipFile(path, 'w', ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def main_case(output, after):
    document = new_document()
    document.settings.odd_and_even_pages_header_footer = True
    section = document.sections[0]
    section.different_first_page_header_footer = True
    project = '演示新路维修工程' if after else '演示旧路维修工程'
    document.add_paragraph('正文项目名称：' + project)
    definitions = {}
    labels = {'header-default': '页眉默认：' + project,
              'header-first': '页眉首页：演示旧路维修工程',
              'header-even': '页眉偶数页：' + project,
              'footer-default': '页脚默认：演示旧路维修工程',
              'footer-first': '页脚首页：' + project,
              'footer-even': '页脚偶数页：' + project}
    for kind in ['header', 'footer']:
        for variant, prefix in [('default', ''), ('first', 'first_page_'), ('even', 'even_page_')]:
            key = kind + '-' + variant
            definitions[key] = fill(getattr(section, prefix + kind), labels[key],
                                    ('页眉表格编号 H-001' if kind == 'header' else '页脚表格编号 F-001') if variant == 'default' else None,
                                    page_field=key == 'footer-default')
    second = document.add_section(WD_SECTION_START.CONTINUOUS)
    second.different_first_page_header_footer = False
    pg = OxmlElement('w:pgNumType')
    pg.set(qn('w:start'), '2')
    second._sectPr.append(pg)
    document.add_paragraph('第二节正文：验证连续分节与沿用。')
    h2text = '第二节默认页眉：' + project
    second_header = fill(second.header, h2text)
    third = document.add_section(WD_SECTION_START.NEW_PAGE)
    third.different_first_page_header_footer = True
    pg = third._sectPr.find(qn('w:pgNumType'))
    if pg is not None:
        third._sectPr.remove(pg)
    document.add_paragraph('第三节正文：验证不同类型分别沿用。')
    third_text = '第三节首页页眉：演示旧路维修工程'
    third_first = fill(third.first_page_header, third_text)
    orphan = Part(PackURI('/word/unused-header.xml'), 'application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml',
                  ('<w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:p><w:r><w:t>游离页眉旧名字</w:t></w:r></w:p></w:hdr>').encode(), document.part.package)
    document.part.relate_to(orphan, RT.HEADER)
    original_default = definitions['header-default']
    if after:
        definitions['header-default'] = 'word/custom-header-A.xml'
    filename = 'headers-after.docx' if after else 'headers-before.docx'
    write_package(output / filename, document, (original_default, definitions['header-default']) if after else None)
    expectations = []
    for key, part in definitions.items():
        expectations.append({'part': part, 'text': labels[key]})
    expectations += [{'part': second_header, 'text': h2text}, {'part': third_first, 'text': third_text},
                     {'part': definitions['header-default'], 'text': '页眉表格编号 H-001', 'table': {'tableIndex': 1, 'xmlRow': 1, 'gridColumn': 2}},
                     {'part': definitions['footer-default'], 'text': '页脚表格编号 F-001', 'table': {'tableIndex': 1, 'xmlRow': 1, 'gridColumn': 2}},
                     {'part': definitions['footer-default'], 'text': '页码：7', 'flag': 'cached_field_result_only'}]
    bindings = []
    for index in [1, 2, 3]:
        for kind in ['header', 'footer']:
            for variant in ['default', 'first', 'even']:
                key = kind + '-' + variant
                part, defined = definitions[key], 1
                if key == 'header-default' and index >= 2:
                    part, defined = second_header, 2
                if key == 'header-first' and index == 3:
                    part, defined = third_first, 3
                bindings.append({'section': index, 'kind': kind, 'variant': variant, 'part': part,
                                 'definedInSection': defined, 'inherited': defined != index,
                                 'enabledBySettings': variant != 'first' or index != 2})
    return {'sectionCount': 3, 'expectations': expectations, 'bindings': bindings,
            'unreferencedParts': ['word/unused-header.xml'], 'duplicateQuote': '重复标记：待核对。',
            'duplicateCount': 8, 'inheritedDefaultFooterPart': definitions['footer-default'],
            'firstHeaderPart': definitions['header-first'], 'defaultHeaderPart': definitions['header-default']}


def switches_case(output):
    document = new_document()
    document.settings.odd_and_even_pages_header_footer = False
    even = document.settings.element.find(qn('w:evenAndOddHeaders'))
    if even is None:
        even = OxmlElement('w:evenAndOddHeaders')
        document.settings.element.append(even)
    even.set(qn('w:val'), '0')
    section = document.sections[0]
    section.different_first_page_header_footer = True
    h = fill(section.header, '开关样例默认页眉')
    f = fill(section.footer, '开关样例默认页脚')
    eh = fill(section.even_page_header, '未启用偶数页页眉：旧项目')
    ef = fill(section.even_page_footer, '未启用偶数页页脚：旧项目')
    document.add_paragraph('第一节首页引用缺失，应记录为空定义。')
    second = document.add_section(WD_SECTION_START.NEW_PAGE)
    second.different_first_page_header_footer = False
    title = OxmlElement('w:titlePg')
    title.set(qn('w:val'), '0')
    second._sectPr.append(title)
    fh = fill(second.first_page_header, '未启用首页页眉：旧项目')
    document.add_paragraph('第二节明确关闭首页不同。')
    write_package(output / 'headers-switches.docx', document)
    return {'sectionCount': 2, 'expectations': [{'part': h, 'text': '开关样例默认页眉'}, {'part': f, 'text': '开关样例默认页脚'},
                                              {'part': eh, 'text': '未启用偶数页页眉：旧项目'}, {'part': ef, 'text': '未启用偶数页页脚：旧项目'},
                                              {'part': fh, 'text': '未启用首页页眉：旧项目'}],
            'bindings': [{'section': 1, 'kind': kind, 'variant': 'first', 'part': None, 'blank': True, 'enabledBySettings': True} for kind in ['header', 'footer']] +
                        [{'section': i, 'kind': kind, 'variant': 'even', 'enabledBySettings': False} for i in [1, 2] for kind in ['header', 'footer']] +
                        [{'section': 2, 'kind': 'header', 'variant': 'first', 'part': fh, 'enabledBySettings': False}], 'unreferencedParts': []}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    if any(args.output.iterdir()):
        raise SystemExit('Use a new empty fixture directory; no files will be overwritten.')
    manifest = {'headers-before.docx': main_case(args.output, False),
                'headers-after.docx': main_case(args.output, True),
                'headers-switches.docx': switches_case(args.output)}
    (args.output / 'header-footer-manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'created': list(manifest), 'synthetic': True}, ensure_ascii=False))
