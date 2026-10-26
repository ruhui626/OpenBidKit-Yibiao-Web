"""Synthetic OOXML fixtures, not customer documents or visually accepted deliverables."""
import argparse
import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_SECTION_START
from lxml import etree

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS = {'w': W, 'mc': 'http://schemas.openxmlformats.org/markup-compatibility/2006',
      'wps': 'http://schemas.microsoft.com/office/word/2010/wordprocessingShape',
      'v': 'urn:schemas-microsoft-com:vml',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
      'future': 'urn:offline-spike:unsupported-future'}


def element(qname, **attrs):
    prefix, local = qname.split(':')
    node = etree.Element(f'{{{NS[prefix]}}}{local}')
    for key, value in attrs.items():
        if ':' in key:
            p, k = key.split(':')
            key = f'{{{NS[p]}}}{k}'
        node.set(key, str(value))
    return node


def paragraph(text):
    p = element('w:p')
    for i, piece in enumerate([text[:3], text[3:]]):
        r = element('w:r')
        if i:
            props = element('w:rPr')
            props.append(element('w:b'))
            r.append(props)
        t = element('w:t')
        t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        t.text = piece
        r.append(t)
        p.append(r)
    return p


class Builder:
    def __init__(self):
        self.doc = Document()
        self.doc.add_heading('文本框提取验证样本', 0)
        self.expected, self.boxes, self.counter = [], [], 0

    def box(self, host, texts, part='word/document.xml', mode='modern', table=False, floating=False, linked=False):
        self.counter += 1
        ident = self.counter
        content = element('w:txbxContent')
        if mode == 'modern':
            outer = element('w:drawing')
            wrap = element('wp:anchor' if floating else 'wp:inline', distT=0, distB=0, distL=0, distR=0)
            if floating:
                for k, v in {'simplePos': '0', 'relativeHeight': '0', 'behindDoc': '0', 'locked': '0', 'layoutInCell': '1', 'allowOverlap': '1'}.items():
                    wrap.set(k, v)
                wrap.append(element('wp:simplePos', x=0, y=0))
                for name in ['positionH', 'positionV']:
                    pos = element('wp:' + name, relativeFrom='column' if name == 'positionH' else 'paragraph')
                    offset = element('wp:posOffset')
                    offset.text = '0'
                    pos.append(offset)
                    wrap.append(pos)
            wrap.append(element('wp:extent', cx=3600000, cy=1000000))
            if floating:
                wrap.append(element('wp:wrapNone'))
            wrap.append(element('wp:docPr', id=ident, name=f'OfflineTextbox{ident}'))
            graphic = element('a:graphic')
            data = element('a:graphicData', uri=NS['wps'])
            shape = element('wps:wsp')
            shape.append(element('wps:cNvSpPr', txBox=1))
            props = element('wps:spPr')
            geom = element('a:prstGeom', prst='rect')
            geom.append(element('a:avLst'))
            props.append(geom)
            shape.append(props)
            txbx = element('wps:linkedTxbx' if linked else 'wps:txbx', id=ident if not linked else 1)
            if linked:
                txbx.set('seq', '1')
            else:
                txbx.append(content)
            shape.append(txbx)
            shape.append(element('wps:bodyPr'))
            data.append(shape)
            graphic.append(data)
            wrap.append(graphic)
            outer.append(wrap)
            key = {'drawingId': str(ident), 'representation': 'drawingml'}
        else:
            outer = element('w:pict')
            shape = element('v:shape', id=f'OfflineVML{ident}', style='width:280pt;height:80pt', coordsize='21600,21600', path='m0,0l21600,0,21600,21600,0,21600xe')
            txbx = element('v:textbox', inset='0,0,0,0')
            txbx.append(content)
            shape.append(txbx)
            outer.append(shape)
            key = {'shapeId': f'OfflineVML{ident}', 'representation': 'vml'}
        host.append(outer)
        if linked:
            return content
        self.boxes.append({'part': part, **key, 'paragraphCount': len(texts)})
        if table:
            tbl = element('w:tbl')
            props = element('w:tblPr')
            props.append(element('w:tblW', **{'w:w': 3000, 'w:type': 'dxa'}))
            tbl.append(props)
            grid = element('w:tblGrid')
            for _ in texts:
                grid.append(element('w:gridCol', **{'w:w': 1500}))
            tbl.append(grid)
            row = element('w:tr')
            for index, text in enumerate(texts, 1):
                cell = element('w:tc')
                cell.append(paragraph(text))
                row.append(cell)
                self.expected.append({'part': part, **key, 'text': text, 'gridColumn': index})
            tbl.append(row)
            content.append(tbl)
            content.append(element('w:p'))
        else:
            for text in texts:
                content.append(paragraph(text))
                if text:
                    self.expected.append({'part': part, **key, 'text': text})
        return content

    def host(self, text='外层锚点'):
        return self.doc.add_paragraph(text).add_run()._r

    def alternate(self, primary, fallback, extra_unknown=False):
        host = self.host('兼容容器锚点')
        group = etree.Element(f'{{{NS["mc"]}}}AlternateContent', nsmap=NS)
        if extra_unknown:
            choice = element('mc:Choice', Requires='future')
            self.box(choice, ['未知分支内容'])
            group.append(choice)
        choice = element('mc:Choice', Requires='wps')
        self.box(choice, [primary])
        group.append(choice)
        fallback_node = element('mc:Fallback')
        self.box(fallback_node, [fallback], mode='legacy')
        group.append(fallback_node)
        host.append(group)


def build_regular(path, changed):
    b = Builder()
    project = '新路' if changed else '旧路'
    b.doc.add_paragraph(f'正文项目：演示{project}维修工程')
    b.box(b.host(), [f'正文文本框：演示{project}维修工程', '同名重复文本'])
    b.box(b.host(), ['正文浮动框：演示旧路维修工程'], floating=True)
    b.box(b.host(), ['旧格式文本框：演示旧路维修工程'], mode='legacy')
    b.box(b.host(), ['同名重复文本'])
    table = b.doc.add_table(rows=1, cols=1)
    host = table.cell(0, 0).paragraphs[0]
    host.add_run('表格单元格锚点')
    b.box(host.add_run()._r, ['框内姓名', '张三'], table=True)
    b.alternate(f'兼容一致：演示{project}维修工程', f'兼容一致：演示{project}维修工程')
    b.alternate(f'兼容差异：演示{project}维修工程', '兼容差异：演示旧路维修工程')
    section = b.doc.sections[0]
    section.different_first_page_header_footer = True
    for story, text, mode in [(section.header, '页眉框：演示旧路维修工程', 'modern'),
                              (section.footer, '页脚框：演示旧路维修工程', 'legacy'),
                              (section.first_page_header, '首页页眉框：演示旧路维修工程', 'modern')]:
        b.box(story.paragraphs[0].add_run()._r, [text], part=str(story.part.partname).lstrip('/'), mode=mode)
    b.doc.add_section(WD_SECTION_START.NEW_PAGE)
    b.doc.add_paragraph('第二节沿用同一页眉页脚')
    b.doc.save(path)
    return {'expected': b.expected, 'boxes': b.boxes, 'expectedGroupsAgree': [True, not changed],
            'unreferencedParts': []}


def build_edges(path):
    b = Builder()
    content = b.box(b.host(), ['前缀😀 工程'])
    p = content[0]
    run = element('w:r')
    run.append(element('w:tab'))
    t = element('w:t'); t.text = '名称'; run.append(t)
    run.append(element('w:br'))
    t = element('w:t'); t.text = '第二行'; run.append(t)
    p.append(run)
    b.expected[-1]['text'] = '前缀😀 工程\t名称\n第二行'
    b.box(b.host(), ['同名/同名'])
    b.box(b.host(), [''])
    content = b.box(b.host(), ['隐藏名称待核对'])
    props = element('w:rPr'); props.append(element('w:vanish')); content[0][0].insert(0, props)
    b.expected[-1]['flag'] = 'hidden_text_requires_review'
    content = b.box(b.host(), ['页码缓存：'])
    field = element('w:fldSimple', **{'w:instr': 'PAGE'})
    r = element('w:r'); t = element('w:t'); t.text = '9'; r.append(t); field.append(r); content[0].append(field)
    b.expected[-1].update(text='页码缓存：9', flag='cached_field_result_only')
    content = b.box(b.host(), ['当前名称：'])
    deleted = element('w:del', **{'w:id': '1', 'w:author': 'Fixture', 'w:date': '2026-10-24T00:00:00Z'})
    r = element('w:r'); t = element('w:delText'); t.text = '删除旧名'; r.append(t); deleted.append(r); content[0].append(deleted)
    added = element('w:ins', **{'w:id': '2', 'w:author': 'Fixture', 'w:date': '2026-10-24T00:00:00Z'})
    r = element('w:r'); t = element('w:t'); t.text = '新增名'; r.append(t); added.append(r); content[0].append(added)
    b.expected[-1].update(text='当前名称：新增名', flag='tracked_changes_require_review')
    b.alternate('新版候选名', '旧版候选名', extra_unknown=True)
    b.box(b.host(), [], linked=True)
    # DrawingML text and VML textpath are deliberately outside the supported Word textbox container.
    host = b.host('未覆盖图形文字')
    pict = element('w:pict'); shape = element('v:shape', id='Watermark', style='width:100pt;height:20pt')
    shape.append(element('v:textpath', string='文本路径水印')); pict.append(shape); host.append(pict)
    drawing = element('w:drawing'); t = element('a:t'); t.text = '图形非文本框文字'; drawing.append(t); host.append(drawing)
    b.doc.save(path)
    # Unreferenced header must be disclosed, never mixed into the used story.
    with ZipFile(path) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    header = element('w:hdr'); p = paragraph('未引用页眉锚点'); header.append(p)
    r = element('w:r'); p.append(r)
    temp = Builder(); temp.box(r, ['未引用页眉文本框'])
    parts['word/unused-textbox-header.xml'] = etree.tostring(header)
    types = etree.fromstring(parts['[Content_Types].xml'])
    etree.SubElement(types, '{http://schemas.openxmlformats.org/package/2006/content-types}Override',
                     PartName='/word/unused-textbox-header.xml', ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml')
    parts['[Content_Types].xml'] = etree.tostring(types)
    with ZipFile(path, 'w', ZIP_DEFLATED) as z:
        for name, raw in parts.items():
            z.writestr(name, raw)
    return {'expected': b.expected, 'boxes': b.boxes, 'expectedGroupsAgree': [False],
            'unreferencedParts': ['word/unused-textbox-header.xml']}


def main():
    parser = argparse.ArgumentParser(); parser.add_argument('--output', type=Path, required=True)
    out = parser.parse_args().output
    out.mkdir(parents=True, exist_ok=False)
    specs = {}
    for name, changed in [('textboxes-before.docx', False), ('textboxes-after.docx', True)]:
        specs[name] = build_regular(out / name, changed)
    specs['textboxes-edges.docx'] = build_edges(out / 'textboxes-edges.docx')
    (out / 'textbox-manifest.json').write_text(json.dumps(specs, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'syntheticFiles': len(specs), 'output': str(out)}))


if __name__ == '__main__':
    main()
